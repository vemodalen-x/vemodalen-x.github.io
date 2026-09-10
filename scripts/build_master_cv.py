from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "outputs" / "cv"
OUTPUT_PATH = OUTPUT_DIR / "Junxian_Wu_AI_Systems_CV_2026.docx"

GRAPHITE = "1F2933"
MUTED = "52606D"
TEAL = "0F766E"
LINE = "D6DEE2"


def set_cell_free_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.50)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)


def set_font(run, size, bold=False, color=GRAPHITE, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_border(paragraph, color=LINE, size="6", space="2"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_hyperlink(paragraph, text, url, color=TEAL):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Aptos")
    r_fonts.set(qn("w:hAnsi"), "Aptos")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color_el, underline])
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.extend([r_pr, text_el])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="CV Section")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text.upper())
    set_font(run, 10.4, bold=True, color=TEAL)
    set_paragraph_border(p)
    return p


def add_body(doc, text, after=2.5, keep=False):
    p = doc.add_paragraph(style="CV Body")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_font(run, 10.0)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="CV Bullet")
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_font(run, 10.0)
    return p


def add_job(doc, title, company, dates, bullets):
    p = doc.add_paragraph(style="CV Job")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(title), 9.7, bold=True)
    set_font(p.add_run(f" | {company}"), 9.7, color=MUTED)
    meta = doc.add_paragraph(style="CV Meta")
    meta.paragraph_format.keep_with_next = True
    set_font(meta.add_run(dates), 8.8, bold=True, color=MUTED)
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_compact_entry(doc, title, meta, text=None):
    p = doc.add_paragraph(style="CV Job")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(title), 9.6, bold=True)
    set_font(p.add_run(f" | {meta}"), 10.0, color=MUTED)
    if text:
        add_body(doc, text, after=3.0)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.0)
    normal.font.color.rgb = RGBColor.from_string(GRAPHITE)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    styles = doc.styles
    for name, base in (
        ("CV Body", "Normal"),
        ("CV Section", "Normal"),
        ("CV Job", "Normal"),
        ("CV Meta", "Normal"),
        ("CV Bullet", "List Bullet"),
    ):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[name]
        style.base_style = styles[base]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")

    styles["CV Body"].paragraph_format.space_after = Pt(2.5)
    styles["CV Body"].paragraph_format.line_spacing = 1.0
    styles["CV Section"].paragraph_format.space_before = Pt(6.5)
    styles["CV Section"].paragraph_format.space_after = Pt(3.0)
    styles["CV Job"].paragraph_format.space_before = Pt(3.2)
    styles["CV Job"].paragraph_format.space_after = Pt(0)
    styles["CV Meta"].paragraph_format.space_after = Pt(1.0)
    bullet = styles["CV Bullet"]
    bullet.paragraph_format.left_indent = Inches(0.18)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_after = Pt(1.35)
    bullet.paragraph_format.line_spacing = 1.0

    # Word's built-in Title style carries a blue template rule by default.
    title = styles["Title"]
    title_paragraph = title._element.get_or_add_pPr()
    title_border = title_paragraph.find(qn("w:pBdr"))
    if title_border is not None:
        title_paragraph.remove(title_border)
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.line_spacing = 1.0


def build_cv():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_cell_free_page(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    configure_styles(doc)

    header = doc.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run("Junxian Wu | Senior AI Engineer"), 8.2, color=MUTED)

    props = doc.core_properties
    props.title = "Junxian Wu - Senior AI Engineer CV"
    props.subject = "Computer Vision, Multimodal and Edge AI Systems"
    props.author = "Junxian Wu"
    props.keywords = "Senior AI Engineer, Computer Vision, Multimodal AI, Edge AI, Forward Deployed Engineer"

    name = doc.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    set_font(name.add_run("JUNXIAN WU"), 23.5, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1.5)
    set_font(title.add_run("Senior AI Engineer | Computer Vision & Edge AI"), 11.4, bold=True, color=TEAL)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(3.5)
    set_font(contact.add_run("Singapore  |  "), 8.3, color=MUTED)
    add_hyperlink(contact, "vemodalenx@gmail.com", "mailto:vemodalenx@gmail.com")
    set_font(contact.add_run("  |  "), 8.3, color=MUTED)
    add_hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/junxian-wu-085aa11b4/")
    set_font(contact.add_run("  |  "), 8.3, color=MUTED)
    add_hyperlink(contact, "GitHub", "https://github.com/vemodalen-x")
    set_font(contact.add_run("  |  "), 8.3, color=MUTED)
    add_hyperlink(contact, "vemodalen-x.github.io", "https://vemodalen-x.github.io/")
    set_paragraph_border(contact, color=TEAL, size="8", space="4")

    add_section_heading(doc, "Profile")
    add_body(
        doc,
        "Senior AI Engineer delivering computer vision and computational photography for mobile camera systems. Experience spans model development, quantization, C++/Android integration, visual-quality evaluation and cross-team release delivery. Independent projects explore agent orchestration and local image-analysis workflows.",
        after=2.0,
    )
    add_section_heading(doc, "Professional Experience")
    add_job(
        doc,
        "Senior AI Engineer",
        "Black Sesame Technologies (Singapore)",
        "Jan 2025 - Present",
        [
            "Own camera AI delivery from product requirements and field feedback through prototypes, hard-case evaluation and release handoff with product, camera tuning, runtime and QA teams.",
            "Lead portrait-imaging development across matting, monocular depth and GPU rendering, investigating subject boundaries, visual artifacts and device behavior.",
            "Develop real-time composition workflows combining object detection, optical-flow tracking, camera guidance and application state.",
            "Drive PyTorch-to-device integration through quantization, TFLite/ONNX export validation, native C++/Android pipelines and on-device debugging.",
        ],
    )
    add_job(
        doc,
        "AI Engineer",
        "Black Sesame Technologies (Singapore)",
        "Aug 2022 - Jan 2025",
        [
            "Productized portrait segmentation, matting and monocular depth models for mobile camera pipelines.",
            "Built training-to-deployment workflows with quantization, export checks and regression comparisons to validate model replacements.",
            "Integrated models into C++/Android runtimes and developed depth-aware bokeh with visual debugging, golden cases and batch release-review tooling.",
        ],
    )

    add_section_heading(doc, "Research Experience")
    add_compact_entry(
        doc,
        "Research Intern",
        "National University of Singapore | Dec 2021 - Jul 2022",
        "Computer-vision research for construction-site safety; co-authored a peer-reviewed Automation in Construction paper.",
    )
    add_compact_entry(
        doc,
        "Research Assistant",
        "Shandong University | Sep 2019 - Jun 2021",
        "Applied computer vision and machine learning to interdisciplinary psychology research, datasets, experiments, and reporting.",
    )

    add_section_heading(doc, "Selected Public Systems")
    public_agent = doc.add_paragraph(style="CV Body")
    set_font(public_agent.add_run("Reliable Agent Workflow: "), 10.0, bold=True, color=TEAL)
    set_font(public_agent.add_run("Scripted-provider Python reference for retries, fallback, output validation and human-review gates, with trace-based regression tests. "), 9.1)
    add_hyperlink(
        public_agent,
        "Code and tests",
        "https://github.com/vemodalen-x/vemodalen-x.github.io/tree/main/agentic_workflow",
    )
    public_mentor = doc.add_paragraph(style="CV Body")
    public_mentor.paragraph_format.space_after = Pt(2)
    set_font(public_mentor.add_run("Photography Mentor: "), 10.0, bold=True, color=TEAL)
    set_font(public_mentor.add_run("Browser-local image diagnostics, source-linked retrieval and practice tracking using deterministic analysis and coaching rules. "), 9.1)
    add_hyperlink(
        public_mentor,
        "Live application",
        "https://vemodalen-x.github.io/photography-mentor-agent.html",
    )

    add_section_heading(doc, "Technical Stack")
    for label, value in [
        ("AI and vision", "PyTorch, TensorFlow, OpenCV, segmentation, matting, monocular depth, detection, optical flow, local retrieval"),
        ("Runtime and quality", "Python, C++, Android/NDK, TFLite, ONNX, TensorRT, OpenGL/GLSL, Docker, quantization, hard cases, traces, release gates"),
    ]:
        p = doc.add_paragraph(style="CV Bullet")
        set_font(p.add_run(f"{label}: "), 9.5, bold=True)
        set_font(p.add_run(value), 9.5)

    add_section_heading(doc, "Education")
    add_compact_entry(doc, "Master of Technology, Intelligent Systems", "National University of Singapore, ISS | 2021 - 2022")
    add_compact_entry(doc, "Bachelor of Engineering, Software Engineering", "Shandong University | 2017 - 2021 | Outstanding Graduate")

    add_section_heading(doc, "Publication")
    pub = doc.add_paragraph(style="CV Body")
    set_font(pub.add_run("Co-author, "), 9.5, bold=True)
    set_font(pub.add_run("\"Automated classification of 'cluttered' construction housekeeping images through supervised and self-supervised feature representation learning,\" "), 9.5)
    set_font(pub.add_run("Automation in Construction"), 9.5, italic=True)
    set_font(pub.add_run(", 2023. "), 9.5)
    add_hyperlink(pub, "doi.org/10.1016/j.autcon.2023.105095", "https://doi.org/10.1016/j.autcon.2023.105095")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_cv()
